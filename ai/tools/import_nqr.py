"""Download the official NQR register and write the engine's snapshot.

Same mechanism as research/03-nqr-import.md: the site's own "Download Summary" button.
No login. Writes data/nqr.json (the fields the engine uses, every row, filtered at
query time) and records the raw file's sha256 so every recommendation can cite it.

    uv run --extra nqr python tools/import_nqr.py
    uv run --extra nqr python tools/import_nqr.py --xlsx saved_copy.xlsx   # offline
"""

import argparse
import datetime as dt
import hashlib
import io
import json
import re
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "data" / "nqr.json"
SEARCH = "https://www.nqr.gov.in/qualifications-search"
DOWNLOAD = "https://www.nqr.gov.in/downloadSummaryFile"
JUNK = {"", "-", "--", "---", "----", "...", "- - -", "n.a.", "na", "n/a", "nil"}


def fetch() -> bytes:
    import httpx

    with httpx.Client(headers={"User-Agent": "Mozilla/5.0"}, timeout=120,
                      follow_redirects=True) as c:
        page = c.get(SEARCH).text
        token = re.search(r'name="_token"\s+value="([^"]+)"', page).group(1)
        ids = re.search(r'name="qualificationids"[^>]*?value="([\d,\s]+)"', page, re.S).group(1)
        ids = ",".join(i.strip() for i in ids.split(",") if i.strip())
        r = c.post(DOWNLOAD, data={"_token": token, "qualificationids": ids},
                   headers={"Referer": SEARCH})
        r.raise_for_status()
        return r.content


def clean(v):
    v = "" if v is None else str(v).strip()
    return None if v.lower() in JUNK else v      # NULL, not guessed


def date(v):
    v = clean(v)
    for fmt in ("%d %b %Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(v, fmt).date().isoformat()
        except (TypeError, ValueError):
            pass
    return None


def hours(v):
    m = re.search(r"\d+", clean(v) or "")
    return int(m.group()) if m else None


def level(v):
    m = re.search(r"\d+(\.\d+)?", clean(v) or "")
    return float(m.group()) if m else None


def parse(xlsx: bytes) -> list[dict]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(xlsx), read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    head_i = next(i for i, r in enumerate(rows) if r and "Title" in r and "Code" in r)
    head = [str(h).strip() if h else "" for h in rows[head_i]]
    out = []
    for r in rows[head_i + 1:]:
        d = dict(zip(head, r))
        if not clean(d.get("Title")):
            continue
        try:
            delivery = json.loads(d.get("Training Delivery Hours") or "{}")
        except json.JSONDecodeError:
            delivery = {}
        out.append({
            "code": clean(d.get("Code")),
            "title": clean(d.get("Title")),
            "description": clean(d.get("Description")),
            "sector": clean(d.get("Sector Name")),
            "level_label": clean(d.get("Level")),        # official string, e.g. "Level 4.5"
            "level": level(d.get("Level")),               # numeric, for comparison only
            "max_hours": hours(d.get("Maximum Notational Hours")),
            "min_hours": hours(d.get("Minimum Notational Hours")),
            "valid_till": date(d.get("Valid Till")),
            "awarding_body": clean(d.get("Awarding Body")),
            "occupation": clean(d.get("Proposed Occupation")),
            "progression": clean(d.get("Progression Pathway")),
            "type": clean(d.get("Qualifcation Type")),   # sic: official header is misspelled
            "ojt_hours": hours(delivery.get("OJT_Mandatory")),
        })
    return out


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--xlsx", type=Path, help="use a saved export instead of downloading")
    a = p.parse_args()
    raw = a.xlsx.read_bytes() if a.xlsx else fetch()
    rows = parse(raw)
    snap = {"source": DOWNLOAD, "fetched": dt.date.today().isoformat(),
            "sha256": hashlib.sha256(raw).hexdigest(), "count": len(rows), "rows": rows}
    OUT.write_text(json.dumps(snap, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    valid = sum(1 for r in rows if r["valid_till"] and r["valid_till"] >= snap["fetched"])
    print(f"{len(rows)} qualifications ({valid} valid today) -> {OUT}  sha256 {snap['sha256'][:12]}")
